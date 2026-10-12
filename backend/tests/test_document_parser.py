"""tests/test_document_parser.py — Comprehensive test suite for Phase 2 Document Parsing & Validation.

Covers:
1. Valid PDF text extraction via PyMuPDF.
2. Valid DOCX text extraction via python-docx (paragraphs and tables).
3. File extension and magic byte verification.
4. Upload size limit enforcement (HTTP 413, FILE_TOO_LARGE).
5. Empty file rejection (HTTP 400).
6. Corrupted PDF & DOCX rejection (HTTP 422, ENCRYPTED_OR_CORRUPT_FILE).
7. Password-protected / encrypted PDF rejection (HTTP 422, ENCRYPTED_OR_CORRUPT_FILE).
8. Text-sparse / empty-text document rejection (HTTP 422, INSUFFICIENT_TEXT_CONTENT).
9. Unsupported extensions and spoofed magic bytes rejection (HTTP 400, INVALID_FILE_TYPE).
10. Page limit enforcement.
11. Privacy test: verifying uploaded resume text is NEVER logged.
12. Error response schema compliance with ErrorResponse envelope.
"""
from __future__ import annotations

import io
import logging
import zipfile
import pytest
import docx
import fitz  # PyMuPDF
from fastapi.testclient import TestClient

from app.config import Settings
from app.exceptions import (
    DocumentValidationError,
    EmptyFileError,
    EncryptedOrCorruptFileError,
    FileTooLargeError,
    InsufficientTextContentError,
    InvalidFileTypeError,
)
from app.main import app
from app.schemas.error import ErrorCode
from app.services.document_parser import (
    clean_extracted_text,
    extract_document_text,
    parse_docx,
    parse_pdf,
)
from app.services.file_validator import validate_file

client = TestClient(app)


# ── Synthetic Document Fixtures ───────────────────────────────────────────────

def make_sample_pdf(text: str = "", pages: int = 1) -> bytes:
    """Create a synthetic valid PDF document in memory."""
    doc = fitz.open()
    for idx in range(pages):
        page = doc.new_page()
        page_text = text or (
            f"Senior Full Stack Engineer Resume - Page {idx + 1}\n"
            "Summary: Experienced software engineer specialized in Python, FastAPI, React, and TypeScript.\n"
            "Experience: Built scalable cloud microservices, high-throughput APIs, and modern web applications.\n"
            "Skills: Python, FastAPI, Next.js, TypeScript, PostgreSQL, Docker, Kubernetes, CI/CD pipelines.\n"
            "Education: Bachelor of Science in Computer Science, University of Technology, 2018-2022.\n"
            "Projects: Developed open-source analytics platform with over 10,000 monthly active users."
        )
        page.insert_text((50, 72), page_text)
    content = doc.write()
    doc.close()
    return content


def make_encrypted_pdf(password: str = "SecretPass123!") -> bytes:
    """Create a synthetic password-protected/encrypted PDF."""
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 72), "Confidential resume data. Should not be parseable without credentials.")
    perm = int(fitz.PDF_PERM_ACCESSIBILITY | fitz.PDF_PERM_PRINT)
    content = doc.write(
        encryption=fitz.PDF_ENCRYPT_AES_256,
        user_pw=password,
        owner_pw=password,
        permissions=perm,
    )
    doc.close()
    return content


def make_sparse_pdf() -> bytes:
    """Create a PDF with almost no text (simulating scanned photo / sparse page)."""
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 72), "Hi")  # Only 1 word, 2 characters
    content = doc.write()
    doc.close()
    return content


def make_sample_docx(
    paragraphs: list[str] | None = None,
    table_data: list[list[str]] | None = None,
) -> bytes:
    """Create a synthetic valid Word DOCX document in memory."""
    doc = docx.Document()
    if paragraphs is None:
        paragraphs = [
            "Senior Backend Engineer — Resume",
            "Professional Summary: Dedicated backend engineer with 6 years of experience designing scalable RESTful APIs.",
            "Work Experience: Designed distributed microservices using Python, FastAPI, Kafka, and Redis.",
            "Key Achievements: Reduced latency by 45% and improved system uptime to 99.99%.",
            "Education: B.Tech in Information Technology, 2016-2020.",
        ]
    for p in paragraphs:
        doc.add_paragraph(p)

    if table_data:
        table = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
        for r_idx, row in enumerate(table_data):
            for c_idx, val in enumerate(row):
                table.cell(r_idx, c_idx).text = val

    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()


def make_sparse_docx() -> bytes:
    """Create a DOCX with fewer than minimum words."""
    doc = docx.Document()
    doc.add_paragraph("Tiny")
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()


# ── 1. Unit Tests: Text Cleaning ──────────────────────────────────────────────

def test_clean_extracted_text_normalizes_crlf_and_controls():
    raw = "Line 1\r\nLine 2\r\x00\x08Line 3   \n\n\n\nLine 4\t"
    cleaned = clean_extracted_text(raw)
    assert "\r" not in cleaned
    assert "\x00" not in cleaned
    assert "Line 1\nLine 2\nLine 3\n\nLine 4" in cleaned


def test_clean_extracted_text_empty_input():
    assert clean_extracted_text("") == ""
    assert clean_extracted_text("   \n\n  ") == ""


# ── 2. Unit Tests: File Validator ─────────────────────────────────────────────

def test_validate_file_valid_pdf():
    pdf_bytes = make_sample_pdf()
    file_type = validate_file(pdf_bytes, "candidate_resume.pdf", 5 * 1024 * 1024, 5)
    assert file_type == "pdf"


def test_validate_file_valid_docx():
    docx_bytes = make_sample_docx()
    file_type = validate_file(docx_bytes, "candidate_resume.docx", 5 * 1024 * 1024, 5)
    assert file_type == "docx"


def test_validate_file_rejects_empty_file():
    with pytest.raises(EmptyFileError) as exc_info:
        validate_file(b"", "empty.pdf", 5 * 1024 * 1024, 5)
    assert exc_info.value.code == ErrorCode.ENCRYPTED_OR_CORRUPT_FILE
    assert "empty" in exc_info.value.message.lower()


def test_validate_file_rejects_oversized():
    huge_bytes = b"x" * (1024 * 1024 + 1)
    with pytest.raises(FileTooLargeError) as exc_info:
        validate_file(huge_bytes, "large.pdf", 1024 * 1024, 1)
    assert exc_info.value.code == ErrorCode.FILE_TOO_LARGE
    assert exc_info.value.status_code == 413


def test_validate_file_rejects_unsupported_extension():
    with pytest.raises(InvalidFileTypeError) as exc_info:
        validate_file(b"some content", "resume.txt", 5 * 1024 * 1024, 5)
    assert exc_info.value.code == ErrorCode.INVALID_FILE_TYPE
    assert "txt" in exc_info.value.message.lower()


def test_validate_file_rejects_spoofed_pdf_magic_bytes():
    fake_pdf = b"This is just plain text masquerading as a PDF document."
    with pytest.raises(InvalidFileTypeError) as exc_info:
        validate_file(fake_pdf, "spoofed.pdf", 5 * 1024 * 1024, 5)
    assert exc_info.value.code == ErrorCode.INVALID_FILE_TYPE
    assert "pdf format" in exc_info.value.message.lower()


def test_validate_file_rejects_spoofed_docx_magic_bytes():
    fake_docx = b"This is not a zip file at all."
    with pytest.raises(InvalidFileTypeError) as exc_info:
        validate_file(fake_docx, "spoofed.docx", 5 * 1024 * 1024, 5)
    assert exc_info.value.code == ErrorCode.INVALID_FILE_TYPE


def test_validate_file_rejects_non_docx_zip_archive():
    # A valid zip, but without Word document parts
    bio = io.BytesIO()
    with zipfile.ZipFile(bio, "w") as zf:
        zf.writestr("test.txt", "hello world")
    zip_bytes = bio.getvalue()

    with pytest.raises(InvalidFileTypeError) as exc_info:
        validate_file(zip_bytes, "malicious.docx", 5 * 1024 * 1024, 5)
    assert exc_info.value.code == ErrorCode.INVALID_FILE_TYPE


# ── 3. Unit Tests: Document Parser Service ────────────────────────────────────

def test_parse_pdf_extracts_clean_text():
    content = make_sample_pdf(pages=2)
    res = parse_pdf(content, "resume.pdf", max_pages=5, min_chars=30, min_words=5)
    assert res.file_type == "pdf"
    assert res.page_count == 2
    assert "FastAPI" in res.raw_text
    assert res.word_count > 10
    assert res.char_count > 50


def test_parse_docx_extracts_paragraphs_and_tables():
    table = [
        ["Category", "Technologies"],
        ["Languages", "Python, TypeScript, Go"],
        ["Frameworks", "FastAPI, Next.js, Django"],
    ]
    content = make_sample_docx(table_data=table)
    res = parse_docx(content, "resume.docx", max_pages=5, min_chars=30, min_words=5)
    assert res.file_type == "docx"
    assert "FastAPI" in res.raw_text
    assert "Languages | Python, TypeScript, Go" in res.raw_text
    assert res.word_count > 10


def test_parse_pdf_rejects_encrypted():
    enc_pdf = make_encrypted_pdf()
    with pytest.raises(EncryptedOrCorruptFileError) as exc_info:
        parse_pdf(enc_pdf, "encrypted.pdf")
    assert exc_info.value.code == ErrorCode.ENCRYPTED_OR_CORRUPT_FILE
    assert "password" in exc_info.value.message.lower() or "encrypted" in exc_info.value.message.lower()


def test_parse_pdf_rejects_corrupted():
    corrupt_pdf = b"%PDF-1.4\nCorrupted content completely invalid %%EOF"
    with pytest.raises(EncryptedOrCorruptFileError) as exc_info:
        parse_pdf(corrupt_pdf, "corrupt.pdf")
    assert exc_info.value.code == ErrorCode.ENCRYPTED_OR_CORRUPT_FILE


def test_parse_pdf_rejects_sparse_text():
    sparse_pdf = make_sparse_pdf()
    with pytest.raises(InsufficientTextContentError) as exc_info:
        parse_pdf(sparse_pdf, "sparse.pdf", min_chars=50, min_words=10)
    assert exc_info.value.code == ErrorCode.INSUFFICIENT_TEXT_CONTENT
    assert "insufficient readable text" in exc_info.value.message.lower()


def test_parse_docx_rejects_sparse_text():
    sparse_docx = make_sparse_docx()
    with pytest.raises(InsufficientTextContentError) as exc_info:
        parse_docx(sparse_docx, "sparse.docx", min_chars=50, min_words=10)
    assert exc_info.value.code == ErrorCode.INSUFFICIENT_TEXT_CONTENT


def test_parse_pdf_rejects_exceeded_pages():
    multi_page_pdf = make_sample_pdf(pages=5)
    with pytest.raises(DocumentValidationError) as exc_info:
        parse_pdf(multi_page_pdf, "long.pdf", max_pages=3)
    assert exc_info.value.status_code == 400
    assert "exceeds" in exc_info.value.message.lower()


# ── 4. API Integration Tests: POST /api/v1/resume/parse ───────────────────────

def test_api_parse_valid_pdf_returns_200():
    pdf_bytes = make_sample_pdf()
    response = client.post(
        "/api/v1/resume/parse",
        files={"file": ("candidate.pdf", pdf_bytes, "application/pdf")},
    )
    assert response.status_code == 200
    data = response.json()
    assert data["filename"] == "candidate.pdf"
    assert data["file_type"] == "pdf"
    assert data["page_count"] == 1
    assert data["word_count"] > 10
    assert "FastAPI" in data["raw_text"]


def test_api_parse_valid_docx_returns_200():
    docx_bytes = make_sample_docx()
    response = client.post(
        "/api/v1/resume/parse",
        files={
            "file": (
                "candidate.docx",
                docx_bytes,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert response.status_code == 200
    data = response.json()
    assert data["filename"] == "candidate.docx"
    assert data["file_type"] == "docx"
    assert data["word_count"] > 10
    assert "RESTful APIs" in data["raw_text"]


def test_api_parse_invalid_extension_returns_400():
    response = client.post(
        "/api/v1/resume/parse",
        files={"file": ("malicious.exe", b"fake binary content", "application/octet-stream")},
    )
    assert response.status_code == 400
    data = response.json()
    assert "error" in data
    assert data["error"]["code"] == ErrorCode.INVALID_FILE_TYPE
    assert "exe" in data["error"]["message"].lower()


def test_api_parse_spoofed_pdf_returns_400():
    response = client.post(
        "/api/v1/resume/parse",
        files={"file": ("fake.pdf", b"plain text without pdf signature", "application/pdf")},
    )
    assert response.status_code == 400
    data = response.json()
    assert data["error"]["code"] == ErrorCode.INVALID_FILE_TYPE
    assert "pdf format" in data["error"]["message"].lower()


def test_api_parse_empty_file_returns_400():
    response = client.post(
        "/api/v1/resume/parse",
        files={"file": ("empty.pdf", b"", "application/pdf")},
    )
    assert response.status_code == 400
    data = response.json()
    assert data["error"]["code"] == ErrorCode.ENCRYPTED_OR_CORRUPT_FILE
    assert "empty" in data["error"]["message"].lower()


def test_api_parse_oversized_file_returns_413():
    from app.config import get_settings

    class SmallLimitSettings(Settings):
        @property
        def max_file_size_bytes(self) -> int:
            return 500

    app.dependency_overrides[get_settings] = lambda: SmallLimitSettings()
    try:
        pdf_bytes = make_sample_pdf()
        assert len(pdf_bytes) > 500

        response = client.post(
            "/api/v1/resume/parse",
            files={"file": ("large.pdf", pdf_bytes, "application/pdf")},
        )
        assert response.status_code == 413
        data = response.json()
        assert data["error"]["code"] == ErrorCode.FILE_TOO_LARGE
    finally:
        app.dependency_overrides.clear()


def test_api_parse_encrypted_pdf_returns_422():
    enc_pdf = make_encrypted_pdf()
    response = client.post(
        "/api/v1/resume/parse",
        files={"file": ("protected.pdf", enc_pdf, "application/pdf")},
    )
    assert response.status_code == 422
    data = response.json()
    assert data["error"]["code"] == ErrorCode.ENCRYPTED_OR_CORRUPT_FILE
    assert "password" in data["error"]["message"].lower() or "encrypted" in data["error"]["message"].lower()


def test_api_parse_corrupt_docx_returns_422():
    # File starts with PK header but has corrupted zip payload
    corrupted_docx = b"PK\x03\x04" + b"\x00" * 200
    response = client.post(
        "/api/v1/resume/parse",
        files={"file": ("corrupt.docx", corrupted_docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert response.status_code == 422
    data = response.json()
    assert data["error"]["code"] == ErrorCode.ENCRYPTED_OR_CORRUPT_FILE


def test_api_parse_sparse_pdf_returns_422():
    sparse_pdf = make_sparse_pdf()
    response = client.post(
        "/api/v1/resume/parse",
        files={"file": ("sparse.pdf", sparse_pdf, "application/pdf")},
    )
    assert response.status_code == 422
    data = response.json()
    assert data["error"]["code"] == ErrorCode.INSUFFICIENT_TEXT_CONTENT


def test_api_parse_sparse_docx_returns_422():
    sparse_docx = make_sparse_docx()
    response = client.post(
        "/api/v1/resume/parse",
        files={"file": ("sparse.docx", sparse_docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert response.status_code == 422
    data = response.json()
    assert data["error"]["code"] == ErrorCode.INSUFFICIENT_TEXT_CONTENT


def test_api_parse_missing_file_parameter_returns_400():
    response = client.post("/api/v1/resume/parse")
    assert response.status_code == 400
    data = response.json()
    assert "error" in data
    assert data["error"]["code"] == ErrorCode.INVALID_FILE_TYPE


# ── 5. Privacy Test: Never Log Resume Contents ────────────────────────────────

def test_logging_never_exposes_resume_content(caplog):
    secret_candidate_marker = "UNIQUE_CANDIDATE_SECRET_MARKER_98765"
    doc_text = (
        f"John Doe Resume — Senior Full Stack Developer\n"
        f"Professional Summary: Architecting scalable Python and TypeScript microservices.\n"
        f"Skills: Python, FastAPI, Docker, Kubernetes, CI/CD, React.\n"
        f"Confidential project code: {secret_candidate_marker}."
    )
    pdf_bytes = make_sample_pdf(text=doc_text)

    with caplog.at_level(logging.DEBUG):
        response = client.post(
            "/api/v1/resume/parse",
            files={"file": ("confidential_resume.pdf", pdf_bytes, "application/pdf")},
        )
    assert response.status_code == 200

    # Ensure secret candidate marker and resume contents are not in ANY log record
    for record in caplog.records:
        assert secret_candidate_marker not in record.message
        assert "John Doe Resume" not in record.message
        assert "Architecting scalable" not in record.message

