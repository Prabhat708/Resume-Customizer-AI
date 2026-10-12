"""app/services/document_parser.py — Document parsing and text extraction engine.

Extracts text from:
- PDF documents using PyMuPDF (fitz).
- DOCX documents using python-docx.

Features:
- Enforces maximum page count.
- Rejects encrypted/password-protected documents.
- Rejects corrupt or unparseable files.
- Cleans and normalizes extracted text.
- Validates text density (rejects empty or text-sparse/scanned documents).
- NEVER logs extracted resume text or candidate personal information.
"""
from __future__ import annotations

import io
import logging
import re
from typing import Optional

import docx
import fitz  # PyMuPDF
from pydantic import BaseModel, Field

from app.config import Settings
from app.exceptions import (
    DocumentValidationError,
    EncryptedOrCorruptFileError,
    InsufficientTextContentError,
    InvalidFileTypeError,
)
from app.schemas.error import ErrorCode

logger = logging.getLogger(__name__)


class DocumentParseResult(BaseModel):
    """Normalized output from document text extraction."""

    filename: str
    file_type: str = Field(..., description="Canonical format: 'pdf' or 'docx'")
    page_count: int
    char_count: int
    word_count: int
    raw_text: str = Field(..., description="Full cleaned extracted text")


def clean_extracted_text(text: str) -> str:
    """Normalize extracted document text.

    - Removes null bytes and non-printable control characters (preserves newlines & tabs).
    - Standardizes CRLF / CR to LF.
    - Strips trailing spaces from individual lines.
    - Collapses excessive blank lines (3+ to 2).
    - Strips leading and trailing outer whitespace.
    """
    if not text:
        return ""

    # Replace carriage returns
    text = text.replace("\r\n", "\n").replace("\r", "\n")

    # Remove null bytes and non-printable control characters
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", "", text)

    # Strip trailing whitespace on each line
    lines = [line.rstrip() for line in text.split("\n")]

    cleaned = "\n".join(lines)

    # Collapse 3 or more consecutive newlines into 2
    cleaned = re.sub(r"\n{3,}", "\n\n", cleaned)

    return cleaned.strip()


def parse_pdf(
    content: bytes,
    filename: str,
    max_pages: int = 10,
    min_chars: int = 50,
    min_words: int = 10,
) -> DocumentParseResult:
    """Extract and validate text from a PDF byte stream using PyMuPDF."""
    doc: Optional[fitz.Document] = None
    try:
        try:
            doc = fitz.open(stream=content, filetype="pdf")
        except Exception as exc:
            err_msg = str(exc).lower()
            if "encrypt" in err_msg or "password" in err_msg:
                raise EncryptedOrCorruptFileError(
                    "The PDF document is password-protected or encrypted. Please upload an unprotected file."
                ) from exc
            raise EncryptedOrCorruptFileError(
                "Failed to read PDF document. The file is corrupted or not a valid PDF."
            ) from exc

        # Check if password-protected / encrypted
        if doc.is_encrypted or doc.needs_pass:
            raise EncryptedOrCorruptFileError(
                "The PDF document is password-protected or encrypted. Please upload an unprotected file."
            )

        page_count = len(doc)
        if page_count == 0:
            raise InsufficientTextContentError("The PDF document contains no pages.")

        if page_count > max_pages:
            raise DocumentValidationError(
                code=ErrorCode.INVALID_FILE_TYPE,
                message=(
                    f"Document contains {page_count} pages, which exceeds the maximum allowed "
                    f"limit of {max_pages} pages."
                ),
                status_code=400,
            )

        extracted_pages = []
        for page_idx in range(page_count):
            try:
                page = doc[page_idx]
                page_text = page.get_text("text")
                if page_text and page_text.strip():
                    extracted_pages.append(page_text.strip())
            except Exception as page_exc:
                logger.warning("Failed extracting page %d: %s", page_idx, str(page_exc))

        combined_text = "\n\n".join(extracted_pages)
        cleaned_text = clean_extracted_text(combined_text)

        word_count = len(cleaned_text.split())
        char_count = len(cleaned_text)

        if char_count < min_chars or word_count < min_words:
            raise InsufficientTextContentError(
                f"The PDF contains insufficient readable text (extracted only {word_count} words). "
                "If this is a scanned document or image-only PDF, please upload a text-based resume."
            )

        logger.info(
            "PDF extraction successful: pages=%d, word_count=%d, char_count=%d",
            page_count,
            word_count,
            char_count,
        )

        return DocumentParseResult(
            filename=filename,
            file_type="pdf",
            page_count=page_count,
            char_count=char_count,
            word_count=word_count,
            raw_text=cleaned_text,
        )

    finally:
        if doc is not None:
            doc.close()


def parse_docx(
    content: bytes,
    filename: str,
    max_pages: int = 10,
    min_chars: int = 50,
    min_words: int = 10,
) -> DocumentParseResult:
    """Extract and validate text from a DOCX byte stream using python-docx."""
    try:
        doc = docx.Document(io.BytesIO(content))
    except Exception as exc:
        err_msg = str(exc).lower()
        if "password" in err_msg or "encrypt" in err_msg:
            raise EncryptedOrCorruptFileError(
                "The DOCX document is password-protected or encrypted. Please upload an unprotected file."
            ) from exc
        raise EncryptedOrCorruptFileError(
            "Failed to open DOCX document. The file is corrupted or not a valid Word document."
        ) from exc

    # Extract paragraphs
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]

    # Extract tables (commonly used for multi-column resumes, skills matrices, headers)
    table_texts = []
    for table in doc.tables:
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells if cell.text and cell.text.strip()]
            # Deduplicate repeated text from merged table cells
            unique_cells = []
            for c in row_cells:
                if not unique_cells or c != unique_cells[-1]:
                    unique_cells.append(c)
            if unique_cells:
                table_texts.append(" | ".join(unique_cells))

    all_parts = paragraphs + table_texts
    combined_text = "\n\n".join(all_parts)
    cleaned_text = clean_extracted_text(combined_text)

    word_count = len(cleaned_text.split())
    char_count = len(cleaned_text)

    if char_count < min_chars or word_count < min_words:
        raise InsufficientTextContentError(
            f"The DOCX contains insufficient readable text (extracted only {word_count} words). "
            "Please upload a document with readable resume text."
        )

    # Word documents are flow documents; estimate page count based on content volume
    estimated_pages = max(1, len(cleaned_text.splitlines()) // 35 + 1)
    if estimated_pages > max_pages:
        raise DocumentValidationError(
            code=ErrorCode.INVALID_FILE_TYPE,
            message=(
                f"Document content exceeds the maximum allowed length (estimated {estimated_pages} pages, "
                f"maximum is {max_pages})."
            ),
            status_code=400,
        )

    logger.info(
        "DOCX extraction successful: estimated_pages=%d, word_count=%d, char_count=%d",
        estimated_pages,
        word_count,
        char_count,
    )

    return DocumentParseResult(
        filename=filename,
        file_type="docx",
        page_count=estimated_pages,
        char_count=char_count,
        word_count=word_count,
        raw_text=cleaned_text,
    )


def extract_document_text(
    content: bytes,
    filename: str,
    file_type: str,
    settings: Settings,
) -> DocumentParseResult:
    """Unified entrypoint for document text extraction based on validated file type."""
    if file_type == "pdf":
        return parse_pdf(
            content=content,
            filename=filename,
            max_pages=settings.max_pages,
            min_chars=settings.min_extracted_chars,
            min_words=settings.min_extracted_words,
        )
    elif file_type == "docx":
        return parse_docx(
            content=content,
            filename=filename,
            max_pages=settings.max_pages,
            min_chars=settings.min_extracted_chars,
            min_words=settings.min_extracted_words,
        )
    else:
        raise InvalidFileTypeError(f"Unsupported file type: {file_type}")
